"""Leitura de modelos de peticao/relatorio em iCloud.

Pasta: ~/Library/Mobile Documents/com~apple~CloudDocs/Modelos TJPA/
Convencao: nome descritivo + .docx ou .md, ex:
  - Modelo peticao contestacao.docx
  - Modelo relatorio despacho geral.docx
  - Modelo manifestacao geral.docx (fallback)
"""
import os
import sys
from pathlib import Path

_base_env = os.environ.get("PJE_STORAGE_DIR")
if _base_env:
    PASTA_MODELOS = Path(_base_env) / "Modelos TJPA"
else:
    PASTA_MODELOS = (
        Path.home()
        / "Library/Mobile Documents/com~apple~CloudDocs/Modelos TJPA"
    )



def _log(msg: str) -> None:
    print(msg, file=sys.stderr, flush=True)


def _tipo_inferido(nome_stem: str) -> str:
    """Heuristica simples pra inferir tipo pelo nome."""
    s = nome_stem.lower()
    if "peticao" in s or "petição" in s or "contestacao" in s or "contestação" in s:
        return "peticao"
    if "relatorio" in s or "relatório" in s:
        return "relatorio"
    if "manifestacao" in s or "manifestação" in s:
        return "manifestacao"
    if "embargo" in s:
        return "embargos"
    if "recurso" in s or "apelacao" in s or "apelação" in s:
        return "recurso"
    return "outro"


def listar_modelos() -> dict:
    """Lista metadados dos modelos da pasta no iCloud."""
    if not PASTA_MODELOS.exists():
        return {
            "pasta": str(PASTA_MODELOS),
            "total": 0,
            "modelos": [],
            "aviso": (
                "Pasta nao existe. Crie manualmente em "
                f"{PASTA_MODELOS} e coloque os modelos .docx/.md la."
            ),
        }

    modelos = []
    for arquivo in sorted(PASTA_MODELOS.iterdir()):
        if arquivo.name.startswith("."):
            continue
        if arquivo.suffix.lower() not in {".docx", ".md", ".txt"}:
            continue
        modelos.append({
            "arquivo": arquivo.name,
            "stem": arquivo.stem,
            "extensao": arquivo.suffix,
            "tipo_inferido": _tipo_inferido(arquivo.stem),
            "tamanho_bytes": arquivo.stat().st_size,
        })

    resultado = {
        "pasta": str(PASTA_MODELOS),
        "total": len(modelos),
        "modelos": modelos,
    }
    if not modelos:
        resultado["aviso"] = (
            "A pasta de modelos existe mas esta VAZIA. Coloque modelos "
            f".docx/.md/.txt em {PASTA_MODELOS} (ex: 'Modelo peticao "
            "contestacao.docx', 'Modelo relatorio analise.docx')."
        )
    return resultado


def ler_modelo(arquivo: str) -> dict:
    """Le o conteudo textual de um modelo.

    arquivo: nome com ou sem extensao. Match case-insensitive.
    """
    if not PASTA_MODELOS.exists():
        raise FileNotFoundError(f"Pasta de modelos nao existe: {PASTA_MODELOS}")

    # Match case-insensitive, com ou sem extensao
    candidatos = []
    arquivo_norm = arquivo.lower().strip()
    for f in PASTA_MODELOS.iterdir():
        if f.name.startswith("."):
            continue
        if f.name.lower() == arquivo_norm:
            candidatos.append(f)
            break
        if f.stem.lower() == arquivo_norm:
            candidatos.append(f)
            break
        # Match parcial
        if arquivo_norm in f.stem.lower():
            candidatos.append(f)

    if not candidatos:
        disponiveis = [
            f.name for f in PASTA_MODELOS.iterdir()
            if f.suffix.lower() in {".docx", ".md", ".txt"} and not f.name.startswith(".")
        ]
        raise FileNotFoundError(
            f"Modelo '{arquivo}' nao encontrado. Disponiveis: {disponiveis}"
        )

    path = candidatos[0]
    _log(f"[MODELO] Lendo {path.name}")

    if path.suffix.lower() == ".docx":
        from docx import Document
        doc = Document(str(path))
        texto = "\n".join(p.text for p in doc.paragraphs)
    else:
        texto = path.read_text(encoding="utf-8")

    return {
        "arquivo": path.name,
        "extensao": path.suffix,
        "tamanho_bytes": path.stat().st_size,
        "conteudo": texto,
    }


def pesquisar_modelos(termo: str) -> dict:
    """Pesquisa termo dentro do nome e conteúdo dos modelos salvos."""
    import re
    if not PASTA_MODELOS.exists():
        return {"termo_busca": termo, "total_encontrados": 0, "modelos": [], "aviso": f"Pasta {PASTA_MODELOS} não existe."}

    termo_norm = termo.lower().strip()
    encontrados = []

    for path in sorted(PASTA_MODELOS.iterdir()):
        if path.name.startswith(".") or path.suffix.lower() not in {".docx", ".md", ".txt"}:
            continue

        nome_match = termo_norm in path.name.lower()
        texto = ""
        try:
            if path.suffix.lower() == ".docx":
                from docx import Document
                doc = Document(str(path))
                texto = "\n".join(p.text for p in doc.paragraphs)
            else:
                texto = path.read_text(encoding="utf-8")
        except Exception:
            pass

        texto_match = termo_norm in texto.lower()

        if nome_match or texto_match:
            snippet = ""
            if texto_match:
                match = re.search(re.escape(termo_norm), texto, re.IGNORECASE)
                if match:
                    st = max(0, match.start() - 80)
                    ed = min(len(texto), match.end() + 80)
                    snippet = texto[st:ed].replace("\n", " ").strip()

            encontrados.append({
                "arquivo": path.name,
                "extensao": path.suffix,
                "match_no_nome": nome_match,
                "match_no_conteudo": texto_match,
                "snippet": f"... {snippet} ..." if snippet else "",
            })

    return {
        "termo_busca": termo,
        "total_encontrados": len(encontrados),
        "modelos": encontrados,
    }
