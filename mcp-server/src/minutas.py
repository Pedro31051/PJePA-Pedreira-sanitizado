"""Salvamento de peticoes e relatorios na pasta do processo (iCloud).

Pasta: ~/Library/Mobile Documents/.../Processos TJPA 1 Grau/{cnj}/
Arquivos: '{tipo} {cnj}.docx' (ex: 'Peticao 0000000-00.0000.0.00.0000.docx')
Se ja existir arquivo com o mesmo nome, versiona: '{tipo} {cnj} (2).docx' etc.
"""
import re
import sys
from pathlib import Path

from pje_downloader import cnj_safe, pasta_processo


def _log(msg: str) -> None:
    print(msg, file=sys.stderr, flush=True)


def _salvar_docx(caminho: Path, conteudo: str) -> None:
    """Salva texto como .docx, paragrafo por paragrafo."""
    from docx import Document
    doc = Document()
    for paragrafo in conteudo.split("\n"):
        doc.add_paragraph(paragrafo)
    doc.save(str(caminho))


def _salvar_texto(caminho: Path, conteudo: str) -> None:
    caminho.write_text(conteudo, encoding="utf-8")


def salvar_peca(
    numero_cnj: str,
    conteudo: str,
    tipo: str = "Petição",
    formato: str = "docx",
    grau: str = "1g",
) -> dict:
    """Salva uma peca (peticao/relatorio/manifestacao/etc) na pasta do processo.

    tipo: 'Petição', 'Relatório', 'Manifestação', 'Despacho' (vai pro nome)
    formato: 'docx' (default) | 'md' | 'txt'
    """
    if formato not in {"docx", "md", "txt"}:
        raise ValueError(f"Formato invalido: {formato}. Use docx/md/txt.")

    pasta = pasta_processo(numero_cnj, grau)
    # Sanitiza tipo e cnj pro nome do arquivo (mesma regra da pasta) -
    # evita separadores de caminho e afins vindos dos parametros
    tipo_seguro = re.sub(r"[^\w À-ÿ-]", "_", tipo).strip() or "Documento"
    base = f"{tipo_seguro} {cnj_safe(numero_cnj)}"
    caminho = pasta / f"{base}.{formato}"

    # Nao sobrescreve versao anterior: acrescenta (2), (3), ...
    versao = 2
    while caminho.exists():
        caminho = pasta / f"{base} ({versao}).{formato}"
        versao += 1

    if formato == "docx":
        _salvar_docx(caminho, conteudo)
    else:
        _salvar_texto(caminho, conteudo)

    tamanho = caminho.stat().st_size
    _log(f"[MINUTA] Salvo {caminho.name} ({tamanho/1024:.1f} KB)")

    return {
        "numero_cnj": numero_cnj,
        "tipo": tipo,
        "formato": formato,
        "caminho": str(caminho),
        "tamanho_bytes": tamanho,
        "tamanho_kb": round(tamanho / 1024, 1),
        "num_caracteres": len(conteudo),
    }


def ler_minuta(numero_cnj: str, arquivo: str, grau: str = "1g", max_caracteres: int = 20000) -> dict:
    """Le de volta uma peca ja gravada na pasta do processo.

    Dava pra salvar e listar minutas, mas nao pra reler: quem quisesse revisar
    ou continuar uma peticao precisava reescreve-la do zero.

    arquivo: nome do arquivo como aparece em listar_minutas_processo.
    """
    pasta = pasta_processo(numero_cnj, grau)
    # Só o nome, nunca um caminho: barra '../' e caminho absoluto.
    nome = Path(arquivo).name
    if not nome:
        return {"erro": "Nome de arquivo vazio.", "numero_cnj": numero_cnj}

    caminho = pasta / nome
    if not caminho.is_file():
        disponiveis = [
            f.name for f in pasta.iterdir()
            if f.is_file() and f.suffix.lower() in {".docx", ".md", ".txt"}
        ] if pasta.exists() else []
        return {
            "erro": f"Minuta não encontrada: '{nome}'",
            "numero_cnj": numero_cnj,
            "pasta": str(pasta),
            "minutas_disponiveis": disponiveis,
        }

    ext = caminho.suffix.lower()
    if ext == ".docx":
        from docx import Document
        doc = Document(str(caminho))
        texto = "\n".join(p.text for p in doc.paragraphs)
    elif ext in {".md", ".txt"}:
        texto = caminho.read_text(encoding="utf-8", errors="replace")
    else:
        return {
            "erro": f"Formato não legível: '{ext}'",
            "formatos_suportados": [".docx", ".md", ".txt"],
            "caminho": str(caminho),
        }

    from datetime import datetime, timezone
    st = caminho.stat()
    return {
        "numero_cnj": numero_cnj,
        "arquivo": caminho.name,
        "caminho": str(caminho),
        "formato": ext.lstrip("."),
        "tamanho_kb": round(st.st_size / 1024, 1),
        "modificado_em": datetime.fromtimestamp(st.st_mtime, tz=timezone.utc).astimezone().strftime("%d/%m/%Y %H:%M:%S"),
        "num_caracteres": len(texto),
        "truncado": len(texto) > max_caracteres,
        "texto": texto[:max_caracteres],
    }


def listar_minutas_processo(numero_cnj: str, grau: str = "1g") -> dict:
    """Lista todas as minutas e pecas salvas (.docx, .md, .txt) na pasta do processo."""
    pasta = pasta_processo(numero_cnj, grau)
    if not pasta.exists():
        return {
            "numero_cnj": numero_cnj,
            "grau": grau,
            "pasta": str(pasta),
            "total_minutas": 0,
            "minutas": [],
        }

    from datetime import datetime, timezone
    minutas = []
    for f in sorted(pasta.iterdir(), key=lambda x: x.stat().st_mtime if x.is_file() else 0, reverse=True):
        if f.is_file() and not f.name.startswith(".") and f.suffix.lower() in {".docx", ".md", ".txt"}:
            mtime = datetime.fromtimestamp(f.stat().st_mtime, tz=timezone.utc).astimezone().strftime("%d/%m/%Y %H:%M:%S")
            minutas.append({
                "arquivo": f.name,
                "caminho": str(f),
                "extensao": f.suffix,
                "tamanho_bytes": f.stat().st_size,
                "tamanho_kb": round(f.stat().st_size / 1024, 1),
                "data_modificacao": mtime,
            })

    return {
        "numero_cnj": numero_cnj,
        "grau": grau,
        "pasta": str(pasta),
        "total_minutas": len(minutas),
        "minutas": minutas,
    }

